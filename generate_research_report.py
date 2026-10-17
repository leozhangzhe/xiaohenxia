from datetime import datetime

from docx import Document


def main() -> None:
    out_path = r"C:\Users\张哲\Desktop\spiderfoot比赛项目\spiderfoot-master\大数据实践赛研究报告-消痕侠.docx"

    doc = Document()
    doc.add_heading("大数据实践赛研究报告", 0)
    p = doc.add_paragraph("项目名称：消痕侠-数字痕迹风险检测与自动处置平台\n")
    p.add_run("赛道方向：大数据实践赛（数据分析、业务智能、辅助决策）\n")
    p.add_run("报告生成时间：" + datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    doc.add_heading("一、研究背景与问题描述", level=1)
    doc.add_paragraph(
        "随着社交平台、公开资料平台和各类互联网服务的普及，个人与组织的数字痕迹持续累积并扩散。"
        "传统人工排查方式存在覆盖面小、效率低、难复检、难量化的问题。"
        "本项目针对“公开信息暴露难发现、发现后难治理、治理后难验证”三类核心痛点，"
        "提出基于多源OSINT数据聚合与风险分析的自动化解决方案。"
    )

    doc.add_heading("二、应用场景", level=1)
    for text in [
        "个人隐私保护：检测邮箱、手机号、用户名、社交账号等公开暴露风险。",
        "组织外部暴露面治理：发现域名、子域名、IP与历史泄露信息的关联风险。",
        "安全运营辅助决策：对风险项按优先级给出处置建议并追踪闭环状态。",
        "竞赛/教学实践：通过可视化看板展示风险趋势、分布与治理成效。",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("三、数据来源与数据特征", level=1)
    doc.add_paragraph(
        "本项目数据主要来自公开互联网情报源（OSINT），包括但不限于域名与证书信息、公开网页内容、"
        "历史泄露线索、网络资产基础信息等。系统通过模块化采集机制将异构数据标准化为统一事件结构，"
        "核心字段包含：事件类型(type)、事件数据(data)、来源模块(module)、风险分值(risk)、"
        "生成时间(generated)、哈希标识(hash)等。"
    )
    doc.add_paragraph(
        "数据特征表现为多源异构、半结构化、噪声较多、更新频繁。为保证可用性，系统在入库阶段进行去重、"
        "关联、规则过滤与风险聚合，形成可计算的风险告警集。"
    )

    doc.add_heading("四、系统设计与开发实现", level=1)
    doc.add_paragraph("系统整体采用“采集层-分析层-处置层-可视化层”四层架构：")
    for text in [
        "采集层：多模块并发采集公开情报，形成扫描事件流。",
        "分析层：基于规则引擎进行关联分析，输出高/中/低风险告警。",
        "处置层：支持一键处理编排（脱敏写回、风险隐藏、复检触发、状态同步）。",
        "可视化层：提供风险告警、数据洞察看板、处置对比报告等界面。",
    ]:
        doc.add_paragraph(text, style="List Number")
    doc.add_paragraph(
        "数据库采用SQLite存储扫描实例、扫描结果、关联结果、处置状态与对比报告。"
        "前端采用Web交互页面展示“扫描概览、风险告警、数据洞察、智能脱敏、一键处理与对比报告”等能力。"
    )

    doc.add_heading("五、数据分析与实验设计", level=1)
    doc.add_paragraph("实验目标：验证系统在“风险发现效率、治理执行能力、治理效果可量化”方面的有效性。")
    doc.add_paragraph("实验流程：")
    for text in [
        "选取目标样本，发起扫描任务并采集风险告警作为基线数据。",
        "执行“一键处理”策略：对敏感数据做掩码写回，对高风险结果执行默认隐藏，并自动触发复检扫描。",
        "复检完成后，自动生成“处理前后对比报告”，计算高风险与中风险变化量。",
        "在看板中呈现风险趋势、风险结构、Top风险类型、处置漏斗与下降率。",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_paragraph("核心评价指标：")
    for text in [
        "高风险下降率 = (处理前高风险数-处理后高风险数)/处理前高风险数",
        "中风险下降率 = (处理前中风险数-处理后中风险数)/处理前中风险数",
        "复检完成率 = 已完成复检报告数/总报告数",
        "处置闭环率 = 已处置并进入待复检或完成复检的风险项占比",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("六、主要结果与结论", level=1)
    doc.add_paragraph("实验结果表明，本项目可实现从“风险发现”到“自动处置”再到“复检对比”的完整闭环：")
    for text in [
        "在风险发现方面：通过多源数据聚合可有效识别公开暴露风险。",
        "在风险治理方面：一键处理可真实执行脱敏写回、默认隐藏和复检触发。",
        "在业务可解释性方面：看板将复杂技术结果转化为普通用户可理解的结论与行动建议。",
        "在量化验证方面：对比报告可直观反映风险变化，支撑治理效果评估。",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("七、创新点", level=1)
    for text in [
        "构建“检测-处置-复检-对比”一体化闭环，不止停留在告警提示。",
        "面向非技术用户提供人话化解释与自动处置建议，降低使用门槛。",
        "引入数据洞察看板和对比报告，强化大数据赛道的分析与决策属性。",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("八、局限性与后续工作", level=1)
    doc.add_paragraph(
        "当前“一键处理”在本系统可控范围内可自动执行；涉及外部平台（社交媒体、搜索引擎）的真实删除与下线"
        "仍依赖平台接口权限或自动化流程。后续将继续扩展外部平台连接器、权限治理机制和跨平台工单编排能力。"
    )

    doc.add_heading("九、可运行程序与实验材料说明", level=1)
    doc.add_paragraph(
        "本项目已提供可运行Web系统、风险告警与数据看板页面、自动处置逻辑和对比报告模块，满足“研究报告+可运行程序+"
        "实验分析+数据支撑”的提交要求。"
    )

    doc.save(out_path)
    print(out_path)


if __name__ == "__main__":
    main()
